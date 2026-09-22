"""文档解析：把 PDF / TXT / DOCX / DOC 抽成纯文本。

- TXT：按 utf-8 -> gb18030 -> utf-16 顺序尝试编码（语文文档常见 gbk/gb18030）
- PDF：pypdf 抽文本（扫描版 PDF 无文本层会得到空串，需 OCR，这里不做）
- DOCX：python-docx，段落 + 表格
- DOC（老格式）：Windows 下用 Word/WPS COM 转换，无 Office 则报错提示
"""
from __future__ import annotations

from pathlib import Path

SUPPORTED_EXTS = {".txt", ".pdf", ".docx", ".doc"}


def is_supported(path: Path | str) -> bool:
    return Path(path).suffix.lower() in SUPPORTED_EXTS


def extract_text(path: Path | str) -> str:
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".txt":
        return _read_txt(p)
    if ext == ".pdf":
        return _read_pdf(p)
    if ext == ".docx":
        return _read_docx(p)
    if ext == ".doc":
        return _read_doc(p)
    raise ValueError(f"不支持的文件格式: {ext}")


def _read_txt(p: Path) -> str:
    data = p.read_bytes()
    for enc in ("utf-8", "gb18030", "utf-16"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="ignore")


def _read_pdf(p: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(p))
    parts = []
    for page in reader.pages:
        txt = page.extract_text() or ""
        if txt.strip():
            parts.append(txt)
    return "\n\n".join(parts)


def _read_docx(p: Path) -> str:
    from docx import Document

    doc = Document(str(p))
    parts = [para.text for para in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _read_doc(p: Path) -> str:
    """老 .doc：借 Word 或 WPS 的 COM 接口转文本。"""
    try:
        import win32com.client  # noqa: F401
    except ImportError as e:
        raise RuntimeError("读取 .doc 需要安装 pywin32（pip install pywin32）") from e

    word = None
    for prog in ("Word.Application", "Kwps.Application", "WPS.Application"):
        try:
            word = win32com.client.Dispatch(prog)
            break
        except Exception:  # noqa: BLE001
            continue
    if word is None:
        raise RuntimeError("读取 .doc 需要本机安装 MS Word 或 WPS（用于 COM 转换）")

    word.Visible = False
    try:
        doc = word.Documents.Open(str(p), ReadOnly=True)
        text = doc.Content.Text
        doc.Close(False)
        return text
    finally:
        try:
            word.Quit()
        except Exception:  # noqa: BLE001
            pass
